#!/usr/bin/env python3
"""Case manager assignment report — case ID and client ID wise (read-only).

Usage (from backend/):
  python -m scripts.export_case_manager_assignments \\
    --output ../exports/insightecase_case_manager_assignments.xlsx

Optional filters:
  --status ACTIVE
  --case-ids 1,2,3
"""
from __future__ import annotations

import argparse
import logging
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.database import SessionLocal  # noqa: E402
from app.models.assignment import CaseAssignment, CaseAssignmentStatus  # noqa: E402
from app.models.case import Case, CaseStatus  # noqa: E402
from app.models.child import Child  # noqa: E402
from app.models.user import User  # noqa: E402

logger = logging.getLogger(__name__)


def _status_val(value: Any) -> str:
    return value.value if hasattr(value, "value") else str(value) if value is not None else ""


def _fmt_date(d: date | None) -> str:
    return d.isoformat() if d else ""


def _fmt_dt(dt: Any) -> str:
    if dt is None:
        return ""
    return dt.isoformat() if hasattr(dt, "isoformat") else str(dt)


def write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 20


def cm_user_fields(user: User | None) -> dict[str, Any]:
    if not user:
        return {
            "case_manager_user_id": "",
            "case_manager_name": "",
            "case_manager_email": "",
            "case_manager_employee_id": "",
            "case_manager_phone": "",
            "case_manager_employment_status": "",
            "case_manager_is_active": "",
        }
    return {
        "case_manager_user_id": user.id,
        "case_manager_name": user.full_name,
        "case_manager_email": user.email,
        "case_manager_employee_id": user.external_employee_id or "",
        "case_manager_phone": user.phone or "",
        "case_manager_employment_status": _status_val(user.employment_status),
        "case_manager_is_active": "yes" if user.is_active else "no",
    }


def run_export(
    *,
    output_path: Path,
    status_filter: str | None = None,
    case_ids: list[int] | None = None,
) -> dict[str, Any]:
    db = SessionLocal()
    try:
        stmt = select(Case).options(selectinload(Case.child)).order_by(Case.id)
        if status_filter:
            try:
                stmt = stmt.where(Case.status == CaseStatus(status_filter))
            except ValueError:
                stmt = stmt.where(Case.status == status_filter)
        if case_ids:
            stmt = stmt.where(Case.id.in_(case_ids))

        cases = list(db.scalars(stmt).all())
        logger.info("Loaded %d cases", len(cases))

        users = {u.id: u for u in db.scalars(select(User)).all()}

        assignments_by_case: dict[int, list[CaseAssignment]] = defaultdict(list)
        if cases:
            case_id_list = [c.id for c in cases]
            for a in db.scalars(select(CaseAssignment).where(CaseAssignment.case_id.in_(case_id_list))).all():
                assignments_by_case[a.case_id].append(a)

        case_rows: list[dict[str, Any]] = []
        missing_cm: list[dict[str, Any]] = []

        for case in cases:
            child = case.child
            cm = users.get(case.case_manager_user_id) if case.case_manager_user_id else None
            cm_fields = cm_user_fields(cm)

            active_asg = [
                a for a in assignments_by_case.get(case.id, [])
                if _status_val(a.status) == CaseAssignmentStatus.ACTIVE.value
            ]
            therapist_ids = sorted({a.therapist_user_id for a in active_asg})
            therapist_names = [
                users[t].full_name for t in therapist_ids if t in users
            ]

            row = {
                "case_id": case.id,
                "case_code": case.case_code,
                "external_case_ref": case.external_case_ref or "",
                "client_id": case.child_id,
                "external_client_id": child.external_client_id if child else "",
                "client_name": child.full_name if child else "",
                **cm_fields,
                "case_manager_assigned": "yes" if case.case_manager_user_id else "no",
                "case_status": _status_val(case.status),
                "status_effective_date": _fmt_date(case.status_effective_date),
                "status_reason": case.status_reason or "",
                "service_category": case.service_type,
                "service_product": case.product_module,
                "region": case.region or "",
                "operational_stage": case.operational_stage or "",
                "billing_type": _status_val(case.billing_type),
                "client_billing_mode": _status_val(case.client_billing_mode),
                "active_therapist_count": len(therapist_ids),
                "active_therapist_ids": "; ".join(str(t) for t in therapist_ids),
                "active_therapist_names": "; ".join(therapist_names),
                "case_created_at": _fmt_dt(case.created_at),
                "case_updated_at": _fmt_dt(case.updated_at),
            }
            case_rows.append(row)
            if not case.case_manager_user_id:
                missing_cm.append(row)

        # Client-wise rollup (one client may have multiple cases)
        by_client: dict[int, list[dict]] = defaultdict(list)
        for row in case_rows:
            by_client[row["client_id"]].append(row)

        client_rows: list[dict[str, Any]] = []
        for client_id in sorted(by_client.keys()):
            rows = by_client[client_id]
            sample = rows[0]
            cm_ids = sorted({r["case_manager_user_id"] for r in rows if r["case_manager_user_id"]})
            cm_names = sorted({r["case_manager_name"] for r in rows if r["case_manager_name"]})
            client_rows.append(
                {
                    "client_id": client_id,
                    "external_client_id": sample["external_client_id"],
                    "client_name": sample["client_name"],
                    "case_count": len(rows),
                    "case_ids": "; ".join(str(r["case_id"]) for r in rows),
                    "case_codes": "; ".join(r["case_code"] for r in rows),
                    "case_manager_user_ids": "; ".join(str(c) for c in cm_ids),
                    "case_manager_names": "; ".join(cm_names),
                    "unique_case_managers": len(cm_ids),
                    "cases_missing_case_manager": sum(1 for r in rows if r["case_manager_assigned"] == "no"),
                    "active_cases": sum(1 for r in rows if r["case_status"] == "ACTIVE"),
                    "service_categories": "; ".join(sorted({r["service_category"] for r in rows if r["service_category"]})),
                }
            )

        # Case manager workload summary
        by_cm: dict[int, list[dict]] = defaultdict(list)
        unassigned_cases: list[dict] = []
        for row in case_rows:
            cm_id = row["case_manager_user_id"]
            if cm_id:
                by_cm[cm_id].append(row)
            else:
                unassigned_cases.append(row)

        cm_summary_rows: list[dict[str, Any]] = []
        for cm_id in sorted(by_cm.keys()):
            rows = by_cm[cm_id]
            sample = rows[0]
            cm_summary_rows.append(
                {
                    "case_manager_user_id": cm_id,
                    "case_manager_name": sample["case_manager_name"],
                    "case_manager_email": sample["case_manager_email"],
                    "case_manager_employee_id": sample["case_manager_employee_id"],
                    "assigned_case_count": len(rows),
                    "active_case_count": sum(1 for r in rows if r["case_status"] == "ACTIVE"),
                    "unique_client_count": len({r["client_id"] for r in rows}),
                    "case_ids": "; ".join(str(r["case_id"]) for r in rows),
                    "client_ids": "; ".join(str(c) for c in sorted({r["client_id"] for r in rows})),
                }
            )

        wb = Workbook()
        wb.remove(wb.active)

        case_headers = list(case_rows[0].keys()) if case_rows else ["case_id"]
        write_sheet(wb, "01_Case_CaseManager", case_headers, case_rows)

        client_headers = list(client_rows[0].keys()) if client_rows else ["client_id"]
        write_sheet(wb, "02_Client_Summary", client_headers, client_rows)

        cm_headers = list(cm_summary_rows[0].keys()) if cm_summary_rows else ["case_manager_user_id"]
        write_sheet(wb, "03_CaseManager_Workload", cm_headers, cm_summary_rows)

        if missing_cm:
            write_sheet(wb, "04_Missing_CaseManager", case_headers, missing_cm)

        summary = [
            {"metric": "Total cases", "value": len(case_rows)},
            {"metric": "Total clients", "value": len(client_rows)},
            {"metric": "Cases with case manager", "value": len(case_rows) - len(missing_cm)},
            {"metric": "Cases missing case manager", "value": len(missing_cm)},
            {"metric": "Unique case managers", "value": len(cm_summary_rows)},
            {"metric": "Active cases", "value": sum(1 for r in case_rows if r["case_status"] == "ACTIVE")},
        ]
        write_sheet(wb, "00_Summary", ["metric", "value"], summary)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        logger.info("Wrote %s (%d cases)", output_path, len(case_rows))

        return {
            "output": str(output_path),
            "cases": len(case_rows),
            "clients": len(client_rows),
            "case_managers": len(cm_summary_rows),
            "missing_cm": len(missing_cm),
        }
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Case manager assignment report (read-only).")
    parser.add_argument("--output", required=True)
    parser.add_argument("--status", dest="status_filter", default=None, help="Filter by case status e.g. ACTIVE")
    parser.add_argument("--case-ids", default=None, help="Comma-separated case IDs")
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO)
    case_ids = [int(x.strip()) for x in args.case_ids.split(",")] if args.case_ids else None
    result = run_export(
        output_path=Path(args.output),
        status_filter=args.status_filter,
        case_ids=case_ids,
    )
    print(
        f"Export complete: {result['output']} | cases={result['cases']} clients={result['clients']} "
        f"case_managers={result['case_managers']} missing_cm={result['missing_cm']}"
    )


if __name__ == "__main__":
    main()
