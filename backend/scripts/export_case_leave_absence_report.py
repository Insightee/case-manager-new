#!/usr/bin/env python3
"""Case-wise leave + child/therapist absence report with reassignment-aware therapist attribution.

Usage (from backend/):
  python -m scripts.export_case_leave_absence_report \\
    --from 2026-07-01 --to 2026-08-01 --timezone Asia/Kolkata \\
    --output ../exports/insightecase_july_2026_leave_absence.xlsx

July 2026 = --from 2026-07-01 --to 2026-08-01 (exclusive end, Asia/Kolkata).
Read-only. Attributes events to therapist assigned on the event date when reassigned.
"""
from __future__ import annotations

import argparse
import json
import logging
import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session, selectinload

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.database import SessionLocal  # noqa: E402
from app.models.assignment import CaseAssignment  # noqa: E402
from app.models.case import Case  # noqa: E402
from app.models.leave import TherapistLeave  # noqa: E402
from app.models.session import Session as TherapySession  # noqa: E402
from app.models.session import SessionStatus  # noqa: E402
from app.models.session_absence import SessionAbsenceRequest, SessionAbsenceType  # noqa: E402
from app.models.user import User  # noqa: E402
from scripts.export_therapist_monthly_sessions import PeriodWindow  # noqa: E402

logger = logging.getLogger(__name__)

IST = ZoneInfo("Asia/Kolkata")
CHILD_ABSENT_STATUSES = {SessionStatus.CLIENT_ABSENT, SessionStatus.NO_SHOW}
THERAPIST_ABSENT_STATUSES = {SessionStatus.THERAPIST_LEAVE}


def _status_val(value: Any) -> str:
    return value.value if hasattr(value, "value") else str(value) if value is not None else ""


def _fmt_date(d: date | None) -> str:
    return d.isoformat() if d else ""


def _fmt_dt(dt: datetime | None) -> str:
    if dt is None:
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ZoneInfo("UTC"))
    return dt.astimezone(IST).isoformat()


def write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 20


def assignments_covering_date(
    assignments: list[CaseAssignment],
    case_id: int,
    on_date: date,
) -> list[CaseAssignment]:
    out: list[CaseAssignment] = []
    for a in assignments:
        if a.case_id != case_id:
            continue
        if a.start_date > on_date:
            continue
        if a.end_date and a.end_date < on_date:
            continue
        out.append(a)
    return out


def primary_assigned_therapist(
    assignments: list[CaseAssignment],
    case_id: int,
    on_date: date,
) -> CaseAssignment | None:
    covering = assignments_covering_date(assignments, case_id, on_date)
    if not covering:
        return None
    active = [a for a in covering if _status_val(a.status) == "ACTIVE"]
    pool = active if active else covering
    return sorted(pool, key=lambda a: (a.start_date, a.id), reverse=True)[0]


def attribution_therapist_id(
    record_therapist_id: int | None,
    case_id: int,
    event_date: date,
    all_assignments: list[CaseAssignment],
) -> tuple[int | None, int | None, str, str]:
    """Return (attributed_therapist_id, assigned_therapist_id, flag, note)."""
    asg = primary_assigned_therapist(all_assignments, case_id, event_date)
    assigned_id = asg.therapist_user_id if asg else None
    if assigned_id is not None:
        attributed = assigned_id
        if record_therapist_id and record_therapist_id != assigned_id:
            return (
                attributed,
                assigned_id,
                "reassignment_adjusted",
                f"Record therapist {record_therapist_id} replaced by assignment {asg.id if asg else ''} on {event_date}",
            )
        return attributed, assigned_id, "matches_assignment", ""
    if record_therapist_id:
        return record_therapist_id, None, "no_assignment_on_date", "Used session/record therapist"
    return None, None, "unassigned", ""


def date_in_window(d: date, window: PeriodWindow) -> bool:
    start = window.start.date()
    end = window.end.date()  # exclusive
    return start <= d < end


def leave_overlaps_window(leave: TherapistLeave, window: PeriodWindow) -> bool:
    w_start = window.start.date()
    w_end = window.end.date() - timedelta(days=1)
    return leave.start_date <= w_end and leave.end_date >= w_start


def expand_leave_case_ids(leave: TherapistLeave) -> list[int]:
    ids: list[int] = []
    if leave.case_id:
        ids.append(leave.case_id)
    raw = leave.case_ids
    if raw:
        if isinstance(raw, list):
            ids.extend(int(x) for x in raw if x is not None)
        elif isinstance(raw, str):
            try:
                parsed = json.loads(raw)
                if isinstance(parsed, list):
                    ids.extend(int(x) for x in parsed if x is not None)
            except json.JSONDecodeError:
                pass
    return sorted(set(ids))


def run_export(
    *,
    from_date: date,
    to_date: date,
    tz_name: str,
    output_path: Path,
) -> dict[str, Any]:
    window = PeriodWindow.from_args(from_date, to_date, tz_name)
    db = SessionLocal()
    try:
        users = {u.id: u for u in db.scalars(select(User)).all()}
        all_assignments = list(db.scalars(select(CaseAssignment)).all())
        cases = {
            c.id: c
            for c in db.scalars(select(Case).options(selectinload(Case.child))).all()
        }

        # Sessions in July window (by scheduled_date buffer + statuses)
        buffer_start = window.start.date() - timedelta(days=1)
        buffer_end = window.end.date() + timedelta(days=1)
        sessions = list(
            db.scalars(
                select(TherapySession)
                .where(
                    TherapySession.scheduled_date >= buffer_start,
                    TherapySession.scheduled_date < buffer_end,
                    or_(
                        TherapySession.status.in_(
                            list(CHILD_ABSENT_STATUSES | THERAPIST_ABSENT_STATUSES)
                        ),
                    ),
                )
                .order_by(TherapySession.case_id, TherapySession.scheduled_date)
            ).all()
        )
        july_absence_sessions = [s for s in sessions if date_in_window(s.scheduled_date, window)]

        absence_requests = list(
            db.scalars(
                select(SessionAbsenceRequest)
                .where(
                    SessionAbsenceRequest.created_at >= window.start,
                    SessionAbsenceRequest.created_at < window.end,
                )
            ).all()
        )
        # Also include requests linked to July sessions
        july_session_ids = {s.id for s in july_absence_sessions}
        for ar in db.scalars(select(SessionAbsenceRequest)).all():
            if ar.session_id in july_session_ids and ar not in absence_requests:
                absence_requests.append(ar)

        leaves = [lv for lv in db.scalars(select(TherapistLeave)).all() if leave_overlaps_window(lv, window)]

        child_absence_rows: list[dict[str, Any]] = []
        therapist_absence_rows: list[dict[str, Any]] = []

        for session in july_absence_sessions:
            case = cases.get(session.case_id)
            if not case:
                continue
            child = case.child
            event_date = session.scheduled_date
            attr_id, asg_id, flag, note = attribution_therapist_id(
                session.therapist_user_id, session.case_id, event_date, all_assignments
            )
            asg = primary_assigned_therapist(all_assignments, session.case_id, event_date)
            therapist = users.get(attr_id) if attr_id else None
            record_therapist = users.get(session.therapist_user_id)

            base = {
                "session_id": session.id,
                "case_id": session.case_id,
                "case_code": case.case_code,
                "client_id": case.child_id,
                "client_name": child.full_name if child else "",
                "event_date": _fmt_date(event_date),
                "session_status": _status_val(session.status),
                "cancellation_reason": session.cancellation_reason or "",
                "record_therapist_user_id": session.therapist_user_id,
                "record_therapist_name": record_therapist.full_name if record_therapist else "",
                "assigned_therapist_user_id": asg_id or "",
                "assigned_therapist_name": users[asg_id].full_name if asg_id and asg_id in users else "",
                "attributed_therapist_user_id": attr_id or "",
                "attributed_therapist_name": therapist.full_name if therapist else "",
                "attributed_employee_id": therapist.external_employee_id if therapist else "",
                "assignment_id": asg.id if asg else "",
                "assignment_start_date": _fmt_date(asg.start_date) if asg else "",
                "assignment_end_date": _fmt_date(asg.end_date) if asg else "",
                "reassignment_reason": asg.reason_for_change if asg else "",
                "attribution_flag": flag,
                "attribution_note": note,
                "service_category": case.service_type,
                "reporting_month": from_date.strftime("%Y-%m"),
            }

            if session.status in CHILD_ABSENT_STATUSES:
                child_absence_rows.append({**base, "absence_type": "CLIENT_ABSENT"})
            if session.status in THERAPIST_ABSENT_STATUSES:
                therapist_absence_rows.append({**base, "absence_type": "THERAPIST_LEAVE_SESSION"})

        # Session absence requests
        request_rows: list[dict[str, Any]] = []
        for ar in absence_requests:
            session = db.get(TherapySession, ar.session_id)
            case = cases.get(ar.case_id)
            if not case:
                continue
            event_date = session.scheduled_date if session else ar.created_at.date()
            if not date_in_window(event_date, window):
                continue
            attr_id, asg_id, flag, note = attribution_therapist_id(
                ar.therapist_user_id, ar.case_id, event_date, all_assignments
            )
            asg = primary_assigned_therapist(all_assignments, ar.case_id, event_date)
            therapist = users.get(attr_id) if attr_id else None
            child = case.child
            request_rows.append(
                {
                    "absence_request_id": ar.id,
                    "session_id": ar.session_id,
                    "case_id": ar.case_id,
                    "case_code": case.case_code,
                    "client_id": case.child_id,
                    "client_name": child.full_name if child else "",
                    "event_date": _fmt_date(event_date),
                    "absence_type": _status_val(ar.absence_type),
                    "request_status": _status_val(ar.status),
                    "reason": ar.reason or "",
                    "notes": ar.notes or "",
                    "billing_outcome": ar.billing_outcome or "",
                    "leave_billing_category": ar.leave_billing_category or "",
                    "therapist_leave_id": ar.therapist_leave_id or "",
                    "record_therapist_user_id": ar.therapist_user_id,
                    "attributed_therapist_user_id": attr_id or "",
                    "attributed_therapist_name": therapist.full_name if therapist else "",
                    "assigned_therapist_user_id": asg_id or "",
                    "assignment_id": asg.id if asg else "",
                    "assignment_start_date": _fmt_date(asg.start_date) if asg else "",
                    "assignment_end_date": _fmt_date(asg.end_date) if asg else "",
                    "reassignment_reason": asg.reason_for_change if asg else "",
                    "attribution_flag": flag,
                    "attribution_note": note,
                    "reviewed_at": _fmt_dt(ar.reviewed_at),
                    "created_at": _fmt_dt(ar.created_at),
                    "reporting_month": from_date.strftime("%Y-%m"),
                }
            )

        # Therapist leave records
        leave_rows: list[dict[str, Any]] = []
        for lv in leaves:
            therapist = users.get(lv.therapist_user_id)
            case_id_list = expand_leave_case_ids(lv)
            if not case_id_list:
                case_id_list = [None]  # type: ignore[list-item]
            for cid in case_id_list:
                case = cases.get(cid) if cid else None
                # Count July days of leave overlapping window
                overlap_start = max(lv.start_date, window.start.date())
                overlap_end = min(lv.end_date, window.end.date() - timedelta(days=1))
                july_days = (overlap_end - overlap_start).days + 1 if overlap_start <= overlap_end else 0
                leave_rows.append(
                    {
                        "leave_id": lv.id,
                        "case_id": cid or "",
                        "case_code": case.case_code if case else "",
                        "client_id": case.child_id if case else "",
                        "client_name": case.child.full_name if case and case.child else "",
                        "therapist_user_id": lv.therapist_user_id,
                        "therapist_name": therapist.full_name if therapist else "",
                        "employee_staff_id": therapist.external_employee_id if therapist else "",
                        "leave_type": _status_val(lv.leave_type),
                        "leave_status": _status_val(lv.status),
                        "billing_category": _status_val(lv.billing_category),
                        "leave_start_date": _fmt_date(lv.start_date),
                        "leave_end_date": _fmt_date(lv.end_date),
                        "july_overlap_days": july_days,
                        "paid_days": lv.paid_days if lv.paid_days is not None else "",
                        "unpaid_days": lv.unpaid_days if lv.unpaid_days is not None else "",
                        "reason": lv.reason or "",
                        "service_line": lv.service_line or "",
                        "consulted_with_parents": "yes" if lv.consulted_with_parents else "no",
                        "includes_shadow_cases": "yes" if lv.includes_shadow_cases else "no",
                        "review_note": lv.review_note or "",
                        "created_at": _fmt_dt(lv.created_at),
                        "reporting_month": from_date.strftime("%Y-%m"),
                    }
                )

        # Case-wise summary (case + attributed therapist)
        case_therapist_stats: dict[tuple[int, int], dict] = defaultdict(
            lambda: {
                "child_absences": 0,
                "therapist_absences": 0,
                "absence_requests": 0,
                "leave_records": 0,
                "july_leave_days": 0,
            }
        )

        for row in child_absence_rows:
            tid = row.get("attributed_therapist_user_id") or row.get("record_therapist_user_id")
            if tid:
                case_therapist_stats[(row["case_id"], int(tid))]["child_absences"] += 1
        for row in therapist_absence_rows:
            tid = row.get("attributed_therapist_user_id") or row.get("record_therapist_user_id")
            if tid:
                case_therapist_stats[(row["case_id"], int(tid))]["therapist_absences"] += 1
        for row in request_rows:
            tid = row.get("attributed_therapist_user_id") or row.get("record_therapist_user_id")
            if tid:
                key = (row["case_id"], int(tid))
                case_therapist_stats[key]["absence_requests"] += 1
                if row["absence_type"] == "CLIENT_ABSENT":
                    case_therapist_stats[key]["child_absences"] += 1
                elif row["absence_type"] == "THERAPIST_LEAVE":
                    case_therapist_stats[key]["therapist_absences"] += 1
        for row in leave_rows:
            if row["case_id"] and row["therapist_user_id"]:
                key = (int(row["case_id"]), int(row["therapist_user_id"]))
                case_therapist_stats[key]["leave_records"] += 1
                case_therapist_stats[key]["july_leave_days"] += int(row["july_overlap_days"] or 0)

        case_summary_rows: list[dict[str, Any]] = []
        for (case_id, therapist_id), stats in sorted(case_therapist_stats.items()):
            case = cases.get(case_id)
            if not case:
                continue
            therapist = users.get(therapist_id)
            asg = primary_assigned_therapist(all_assignments, case_id, window.start.date())
            case_summary_rows.append(
                {
                    "case_id": case_id,
                    "case_code": case.case_code,
                    "client_id": case.child_id,
                    "client_name": case.child.full_name if case.child else "",
                    "therapist_user_id": therapist_id,
                    "therapist_name": therapist.full_name if therapist else "",
                    "employee_staff_id": therapist.external_employee_id if therapist else "",
                    "service_category": case.service_type,
                    "child_absence_sessions_july": stats["child_absences"],
                    "therapist_absence_sessions_july": stats["therapist_absences"],
                    "absence_requests_july": stats["absence_requests"],
                    "leave_records_july": stats["leave_records"],
                    "july_leave_overlap_days": stats["july_leave_days"],
                    "current_assignment_id": asg.id if asg else "",
                    "current_assignment_start": _fmt_date(asg.start_date) if asg else "",
                    "current_assignment_end": _fmt_date(asg.end_date) if asg else "",
                    "reporting_month": from_date.strftime("%Y-%m"),
                }
            )

        # Case rollup (all therapists)
        case_rollup: dict[int, dict] = defaultdict(lambda: {"child": 0, "therapist": 0, "requests": 0, "leave": 0})
        for row in case_summary_rows:
            cid = row["case_id"]
            case_rollup[cid]["child"] += row["child_absence_sessions_july"]
            case_rollup[cid]["therapist"] += row["therapist_absence_sessions_july"]
            case_rollup[cid]["requests"] += row["absence_requests_july"]
            case_rollup[cid]["leave"] += row["leave_records_july"]

        case_only_rows: list[dict[str, Any]] = []
        for case_id in sorted(case_rollup.keys()):
            case = cases[case_id]
            stats = case_rollup[case_id]
            therapists = sorted(
                {r["therapist_user_id"] for r in case_summary_rows if r["case_id"] == case_id}
            )
            case_only_rows.append(
                {
                    "case_id": case_id,
                    "case_code": case.case_code,
                    "client_id": case.child_id,
                    "client_name": case.child.full_name if case.child else "",
                    "therapist_ids_involved": "; ".join(str(t) for t in therapists),
                    "therapist_count": len(therapists),
                    "multiple_therapists_flag": "yes" if len(therapists) > 1 else "no",
                    "child_absence_sessions_july": stats["child"],
                    "therapist_absence_sessions_july": stats["therapist"],
                    "absence_requests_july": stats["requests"],
                    "leave_records_july": stats["leave"],
                    "service_category": case.service_type,
                    "reporting_month": from_date.strftime("%Y-%m"),
                }
            )

        reassignment_rows = [
            r for r in child_absence_rows + therapist_absence_rows + request_rows
            if r.get("attribution_flag") == "reassignment_adjusted"
        ]

        wb = Workbook()
        wb.remove(wb.active)

        summary = [
            {"metric": "Reporting period (IST)", "value": f"{from_date} to {to_date} (exclusive)"},
            {"metric": "Child absence sessions (July)", "value": len(child_absence_rows)},
            {"metric": "Therapist absence sessions (July)", "value": len(therapist_absence_rows)},
            {"metric": "Absence requests (July)", "value": len(request_rows)},
            {"metric": "Leave records overlapping July", "value": len({r['leave_id'] for r in leave_rows})},
            {"metric": "Case-therapist pairs with events", "value": len(case_summary_rows)},
            {"metric": "Cases with any absence/leave", "value": len(case_only_rows)},
            {"metric": "Reassignment-adjusted rows", "value": len(reassignment_rows)},
        ]
        write_sheet(wb, "00_Summary", ["metric", "value"], summary)

        if case_only_rows:
            write_sheet(wb, "01_Case_Summary", list(case_only_rows[0].keys()), case_only_rows)
        if case_summary_rows:
            write_sheet(wb, "02_Case_Therapist_Summary", list(case_summary_rows[0].keys()), case_summary_rows)
        if child_absence_rows:
            write_sheet(wb, "03_Child_Absence_Detail", list(child_absence_rows[0].keys()), child_absence_rows)
        if therapist_absence_rows:
            write_sheet(wb, "04_Therapist_Absence_Detail", list(therapist_absence_rows[0].keys()), therapist_absence_rows)
        if request_rows:
            write_sheet(wb, "05_Absence_Requests", list(request_rows[0].keys()), request_rows)
        if leave_rows:
            write_sheet(wb, "06_Therapist_Leave", list(leave_rows[0].keys()), leave_rows)
        if reassignment_rows:
            headers = list(reassignment_rows[0].keys())
            write_sheet(wb, "07_Reassignment_Adjusted", headers, reassignment_rows)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        logger.info("Wrote %s", output_path)

        return {
            "output": str(output_path),
            "child_absences": len(child_absence_rows),
            "therapist_absences": len(therapist_absence_rows),
            "requests": len(request_rows),
            "leaves": len({r["leave_id"] for r in leave_rows}),
            "cases": len(case_only_rows),
            "reassignment_adjusted": len(reassignment_rows),
        }
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="July leave/absence case report (read-only).")
    parser.add_argument("--from", dest="from_date", default="2026-07-01")
    parser.add_argument("--to", dest="to_date", default="2026-08-01")
    parser.add_argument("--timezone", default="Asia/Kolkata")
    parser.add_argument("--output", required=True)
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO)
    result = run_export(
        from_date=date.fromisoformat(args.from_date),
        to_date=date.fromisoformat(args.to_date),
        tz_name=args.timezone,
        output_path=Path(args.output),
    )
    print(
        f"Export: {result['output']} | child_abs={result['child_absences']} "
        f"therapist_abs={result['therapist_absences']} requests={result['requests']} "
        f"leaves={result['leaves']} cases={result['cases']} reassigned={result['reassignment_adjusted']}"
    )


if __name__ == "__main__":
    main()
