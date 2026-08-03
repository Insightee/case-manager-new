#!/usr/bin/env python3
"""Case-wise client billing + therapist payout report from attendance (not log approval).

Usage (from backend/):
  python -m scripts.export_case_attendance_billing_report \\
    --from 2026-07-01 --to 2026-08-01 --timezone Asia/Kolkata \\
    --output ../exports/insightecase_july_2026_case_billing_by_attendance.xlsx

Counts sessions with check-in/out attendance evidence in the IST window.
Does NOT require daily_log approval. Read-only — no DB writes.
"""
from __future__ import annotations

import argparse
import logging
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session, selectinload

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.database import SessionLocal  # noqa: E402
from app.models.assignment import CaseAssignment  # noqa: E402
from app.models.case import BillingType, Case, CompensationMode  # noqa: E402
from app.models.daily_log import LogApprovalStatus  # noqa: E402
from app.models.invoice_line import SessionLineType  # noqa: E402
from app.models.session import Session as TherapySession  # noqa: E402
from app.models.session import SessionStatus  # noqa: E402
from app.models.user import User  # noqa: E402
from app.services.invoice_billing_service import compute_session_line_amount  # noqa: E402
from scripts.export_therapist_monthly_sessions import (  # noqa: E402
    LogSnapshot,
    PeriodWindow,
    duration_minutes,
    effective_times_for_export,
    fetch_log_snapshots,
    resolve_reporting_start,
)

logger = logging.getLogger(__name__)

IST = ZoneInfo("Asia/Kolkata")
SKIP_STATUSES = {
    SessionStatus.CANCELLED,
    SessionStatus.RESCHEDULED,
    SessionStatus.NO_SHOW,
    SessionStatus.THERAPIST_LEAVE,
}


def _status_val(value: Any) -> str:
    return value.value if hasattr(value, "value") else str(value) if value is not None else ""


def _safe_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _fmt_date(d: date | None) -> str:
    return d.isoformat() if d else ""


@dataclass
class AttendedSession:
    session_id: int
    case_id: int
    therapist_user_id: int
    reporting_date: date
    duration_minutes: int
    log_approved: bool
    client_absent: bool
    session_status: str


def is_attendance_occurred(
    session: TherapySession,
    log: LogSnapshot | None,
    window: PeriodWindow,
) -> tuple[bool, date | None]:
    """Attendance-based occurred — does NOT require daily log approval."""
    if session.status in SKIP_STATUSES:
        return False, None

    reporting_start, _ = resolve_reporting_start(session, log)
    if reporting_start is None:
        return False, None
    if not (window.start <= reporting_start < window.end):
        return False, None

    has_start = bool(session.actual_start_at)
    has_end = bool(session.actual_end_at) or bool(session.auto_ended)

    if has_start and (has_end or session.auto_ended):
        return True, reporting_start.date()

    if has_start and session.status == SessionStatus.COMPLETED:
        return True, reporting_start.date()

    if log and log.attendance_status in ("PRESENT", "LATE", "PARTIAL") and has_start:
        return True, reporting_start.date()

    return False, None


def client_rate_per_session(case: Case) -> float | None:
    if case.billing_type == BillingType.PER_SESSION:
        return _safe_float(case.client_rate_per_session_inr)
    pkg = int(case.package_session_count or 0)
    if pkg and case.package_amount_inr:
        return round(float(case.package_amount_inr) / pkg, 2)
    return _safe_float(case.client_rate_per_session_inr)


def payout_for_session_index(case: Case, index: int) -> float:
    if case.billing_type == BillingType.PER_SESSION:
        line_type = SessionLineType.PER_SESSION
    else:
        pkg = int(case.package_session_count or 0)
        line_type = SessionLineType.INCLUDED if index < pkg else SessionLineType.ADDITIONAL
    return compute_session_line_amount(case, line_type)


def fetch_sessions(db: Session, window: PeriodWindow) -> list[TherapySession]:
    buffer_start = window.start.date() - timedelta(days=3)
    buffer_end = window.end.date() + timedelta(days=3)
    stmt = (
        select(TherapySession)
        .where(
            or_(
                and_(
                    TherapySession.scheduled_date >= buffer_start,
                    TherapySession.scheduled_date <= buffer_end,
                ),
                TherapySession.actual_start_at.isnot(None),
            )
        )
        .options(selectinload(TherapySession.case).selectinload(Case.child))
        .order_by(TherapySession.case_id, TherapySession.scheduled_date, TherapySession.id)
    )
    return list(db.scalars(stmt).all())


def assignment_for_therapist(
    assignments: list[CaseAssignment],
    therapist_user_id: int,
) -> CaseAssignment | None:
    matching = [a for a in assignments if a.therapist_user_id == therapist_user_id]
    if not matching:
        return None
    active = [a for a in matching if _status_val(a.status) == "ACTIVE"]
    return active[0] if active else sorted(matching, key=lambda a: a.start_date, reverse=True)[0]


def build_assignment_history_row(
    case: Case,
    child_name: str,
    assignment: CaseAssignment,
    users: dict[int, User],
) -> dict[str, Any]:
    therapist = users.get(assignment.therapist_user_id)
    assigned_by = users.get(assignment.assigned_by_user_id) if assignment.assigned_by_user_id else None
    return {
        "case_id": case.id,
        "case_code": case.case_code,
        "client_id": case.child_id,
        "client_name": child_name,
        "assignment_id": assignment.id,
        "therapist_user_id": assignment.therapist_user_id,
        "therapist_name": therapist.full_name if therapist else "",
        "employee_staff_id": therapist.external_employee_id if therapist else "",
        "assignment_start_date": _fmt_date(assignment.start_date),
        "assignment_end_date": _fmt_date(assignment.end_date),
        "assignment_status": _status_val(assignment.status),
        "reassignment_reason": assignment.reason_for_change or "",
        "assignment_notes": assignment.notes or "",
        "assigned_by": assigned_by.full_name if assigned_by else "",
        "case_status": _status_val(case.status),
        "service_category": case.service_type,
    }


def write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18


def run_report(
    *,
    from_date: date,
    to_date: date,
    tz_name: str,
    output_path: Path,
) -> dict[str, Any]:
    window = PeriodWindow.from_args(from_date, to_date, tz_name)
    db = SessionLocal()
    try:
        sessions = fetch_sessions(db, window)
        session_ids = [s.id for s in sessions]
        logs = fetch_log_snapshots(db, session_ids)

        users = {u.id: u for u in db.scalars(select(User)).all()}
        case_ids = list({s.case_id for s in sessions})
        assignments_by_case: dict[int, list[CaseAssignment]] = defaultdict(list)
        for a in db.scalars(select(CaseAssignment).where(CaseAssignment.case_id.in_(case_ids))).all():
            assignments_by_case[a.case_id].append(a)
        for cid in assignments_by_case:
            assignments_by_case[cid].sort(key=lambda x: (x.start_date, x.id))

        attended: list[AttendedSession] = []
        for session in sessions:
            log = logs.get(session.id)
            ok, rep_date = is_attendance_occurred(session, log, window)
            if not ok or rep_date is None:
                continue
            eff_start, eff_end = effective_times_for_export(session, log)
            dur = duration_minutes(session, log, eff_start, eff_end) or 0
            attended.append(
                AttendedSession(
                    session_id=session.id,
                    case_id=session.case_id,
                    therapist_user_id=session.therapist_user_id,
                    reporting_date=rep_date,
                    duration_minutes=dur,
                    log_approved=bool(log and log.approval_status == LogApprovalStatus.APPROVED.value),
                    client_absent=session.status == SessionStatus.CLIENT_ABSENT,
                    session_status=_status_val(session.status),
                )
            )

        cases = {c.id: c for c in db.scalars(
            select(Case).where(Case.id.in_(case_ids)).options(selectinload(Case.child))
        ).all()}

        # Global session index per case for package payout ordering
        by_case_sessions: dict[int, list[AttendedSession]] = defaultdict(list)
        for a in attended:
            by_case_sessions[a.case_id].append(a)
        for cid in by_case_sessions:
            by_case_sessions[cid].sort(key=lambda x: (x.reporting_date, x.session_id))

        session_payout: dict[int, float] = {}
        for cid, sess_list in by_case_sessions.items():
            case = cases.get(cid)
            if not case or not case.billing_type:
                continue
            for idx, s in enumerate(sess_list):
                session_payout[s.session_id] = payout_for_session_index(case, idx)

        # Aggregate case + therapist
        ct_key_sessions: dict[tuple[int, int], list[AttendedSession]] = defaultdict(list)
        for a in attended:
            ct_key_sessions[(a.case_id, a.therapist_user_id)].append(a)

        case_therapist_rows: list[dict[str, Any]] = []
        for (case_id, therapist_id), sess_list in sorted(ct_key_sessions.items()):
            case = cases.get(case_id)
            if not case:
                continue
            child = case.child
            therapist = users.get(therapist_id)
            asg_list = assignments_by_case.get(case_id, [])
            asg = assignment_for_therapist(asg_list, therapist_id)
            other_therapists = sorted(
                {a.therapist_user_id for a in asg_list if a.therapist_user_id != therapist_id}
            )
            other_names = [
                users[t].full_name for t in other_therapists if t in users
            ]

            attended_count = len(sess_list)
            approved_count = sum(1 for s in sess_list if s.log_approved)
            client_absent_count = sum(1 for s in sess_list if s.client_absent)
            total_mins = sum(s.duration_minutes for s in sess_list)
            rate = client_rate_per_session(case)
            client_total = round((rate or 0) * attended_count, 2) if rate else ""
            payout_total = round(sum(session_payout.get(s.session_id, 0) for s in sess_list), 2)

            per_sess_payout = ""
            if case.billing_type == BillingType.PER_SESSION:
                per_sess_payout = _safe_float(case.pay_share_amount_inr or case.therapist_fixed_pay_inr)
            elif case.package_session_count:
                per_sess_payout = payout_for_session_index(case, 0)

            case_therapist_rows.append(
                {
                    "case_id": case_id,
                    "case_code": case.case_code,
                    "external_case_ref": case.external_case_ref or "",
                    "client_id": case.child_id,
                    "external_client_id": child.external_client_id if child else "",
                    "client_name": child.full_name if child else "",
                    "therapist_user_id": therapist_id,
                    "employee_staff_id": therapist.external_employee_id if therapist else "",
                    "therapist_name": therapist.full_name if therapist else "",
                    "assignment_id": asg.id if asg else "",
                    "assignment_start_date": _fmt_date(asg.start_date) if asg else "",
                    "assignment_end_date": _fmt_date(asg.end_date) if asg else "",
                    "assignment_status": _status_val(asg.status) if asg else "",
                    "reassignment_reason": asg.reason_for_change if asg else "",
                    "other_therapists_on_case": "; ".join(other_names),
                    "other_therapist_ids": "; ".join(str(t) for t in other_therapists),
                    "multiple_therapists_flag": "yes" if other_therapists else "no",
                    "case_status": _status_val(case.status),
                    "case_start_effective": _fmt_date(case.status_effective_date),
                    "service_category": case.service_type,
                    "service_product": case.product_module,
                    "billing_type": _status_val(case.billing_type),
                    "compensation_mode": _status_val(case.compensation_mode),
                    "client_billing_mode": _status_val(case.client_billing_mode),
                    "client_rate_per_session_inr": _safe_float(case.client_rate_per_session_inr) or "",
                    "package_amount_inr": _safe_float(case.package_amount_inr) or "",
                    "package_session_count": case.package_session_count or "",
                    "pay_share_amount_inr": _safe_float(case.pay_share_amount_inr) or "",
                    "therapist_fixed_pay_inr": _safe_float(case.therapist_fixed_pay_inr) or "",
                    "attended_sessions_count": attended_count,
                    "approved_log_sessions_count": approved_count,
                    "attendance_minus_approved_logs": attended_count - approved_count,
                    "client_absent_sessions_count": client_absent_count,
                    "total_attended_minutes": total_mins,
                    "total_attended_hours": round(total_mins / 60, 2),
                    "calculated_client_billing_inr": client_total,
                    "calculated_therapist_payout_inr": payout_total,
                    "avg_payout_per_attended_session": round(payout_total / attended_count, 2) if attended_count else "",
                    "billing_notes": case.billing_notes or "",
                    "calculation_basis": "attendance_timestamps_not_log_approval",
                    "reporting_month": from_date.strftime("%Y-%m"),
                }
            )

        # Case summary (rollup across therapists)
        case_summary_rows: list[dict[str, Any]] = []
        by_case: dict[int, list[dict]] = defaultdict(list)
        for row in case_therapist_rows:
            by_case[row["case_id"]].append(row)

        for case_id, rows in sorted(by_case.items()):
            case = cases.get(case_id)
            if not case:
                continue
            child = case.child
            therapists = sorted({r["therapist_name"] for r in rows if r["therapist_name"]})
            therapist_ids = sorted({r["therapist_user_id"] for r in rows})
            case_summary_rows.append(
                {
                    "case_id": case_id,
                    "case_code": case.case_code,
                    "client_id": case.child_id,
                    "external_client_id": child.external_client_id if child else "",
                    "client_name": child.full_name if child else "",
                    "therapist_count": len(therapist_ids),
                    "therapist_ids": "; ".join(str(t) for t in therapist_ids),
                    "therapist_names": "; ".join(therapists),
                    "multiple_therapists_flag": "yes" if len(therapist_ids) > 1 else "no",
                    "service_category": case.service_type,
                    "billing_type": _status_val(case.billing_type),
                    "client_rate_per_session_inr": _safe_float(case.client_rate_per_session_inr) or "",
                    "total_attended_sessions": sum(r["attended_sessions_count"] for r in rows),
                    "total_approved_log_sessions": sum(r["approved_log_sessions_count"] for r in rows),
                    "attendance_minus_approved": sum(r["attendance_minus_approved_logs"] for r in rows),
                    "total_client_billing_inr": sum(float(r["calculated_client_billing_inr"] or 0) for r in rows),
                    "total_therapist_payout_inr": sum(float(r["calculated_therapist_payout_inr"] or 0) for r in rows),
                    "reporting_month": from_date.strftime("%Y-%m"),
                }
            )

        # Assignment history for all cases with activity
        active_case_ids = set(by_case.keys())
        assignment_rows: list[dict[str, Any]] = []
        for case_id in sorted(active_case_ids):
            case = cases.get(case_id)
            if not case:
                continue
            child_name = case.child.full_name if case.child else ""
            for asg in assignments_by_case.get(case_id, []):
                assignment_rows.append(build_assignment_history_row(case, child_name, asg, users))

        # Session detail
        detail_rows: list[dict[str, Any]] = []
        for a in sorted(attended, key=lambda x: (x.case_id, x.reporting_date, x.session_id)):
            case = cases.get(a.case_id)
            therapist = users.get(a.therapist_user_id)
            rate = client_rate_per_session(case) if case else None
            detail_rows.append(
                {
                    "session_id": a.session_id,
                    "case_id": a.case_id,
                    "case_code": case.case_code if case else "",
                    "client_id": case.child_id if case else "",
                    "client_name": case.child.full_name if case and case.child else "",
                    "therapist_user_id": a.therapist_user_id,
                    "therapist_name": therapist.full_name if therapist else "",
                    "reporting_date_ist": a.reporting_date.isoformat(),
                    "session_status": a.session_status,
                    "duration_minutes": a.duration_minutes,
                    "log_approved": "yes" if a.log_approved else "no",
                    "client_absent": "yes" if a.client_absent else "no",
                    "client_billing_inr": rate if rate else "",
                    "therapist_payout_inr": session_payout.get(a.session_id, ""),
                    "included_by_attendance_not_log": "yes",
                }
            )

        wb = Workbook()
        wb.remove(wb.active)

        ct_headers = list(case_therapist_rows[0].keys()) if case_therapist_rows else ["case_id"]
        write_sheet(wb, "01_Case_Therapist_Billing", ct_headers, case_therapist_rows)

        cs_headers = list(case_summary_rows[0].keys()) if case_summary_rows else ["case_id"]
        write_sheet(wb, "02_Case_Summary", cs_headers, case_summary_rows)

        asg_headers = list(assignment_rows[0].keys()) if assignment_rows else ["case_id"]
        write_sheet(wb, "03_Assignment_History", asg_headers, assignment_rows)

        det_headers = list(detail_rows[0].keys()) if detail_rows else ["session_id"]
        write_sheet(wb, "04_Session_Detail", det_headers, detail_rows)

        # Totals sheet
        totals = [
            {"metric": "Cases with attendance", "value": len(case_summary_rows)},
            {"metric": "Case-therapist pairs", "value": len(case_therapist_rows)},
            {"metric": "Total attended sessions", "value": len(attended)},
            {"metric": "Sessions with approved logs", "value": sum(1 for a in attended if a.log_approved)},
            {"metric": "Attendance minus approved logs", "value": len(attended) - sum(1 for a in attended if a.log_approved)},
            {"metric": "Total client billing INR", "value": round(sum(float(r["total_client_billing_inr"]) for r in case_summary_rows), 2)},
            {"metric": "Total therapist payout INR", "value": round(sum(float(r["total_therapist_payout_inr"]) for r in case_summary_rows), 2)},
            {"metric": "Cases with multiple therapists", "value": sum(1 for r in case_summary_rows if r["multiple_therapists_flag"] == "yes")},
        ]
        write_sheet(wb, "00_Summary", ["metric", "value"], totals)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        logger.info("Wrote %s (%d case-therapist rows)", output_path, len(case_therapist_rows))

        return {
            "output": str(output_path),
            "case_therapist_rows": len(case_therapist_rows),
            "cases": len(case_summary_rows),
            "attended_sessions": len(attended),
            "totals": totals,
        }
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Case-wise attendance billing report (read-only).")
    parser.add_argument("--from", dest="from_date", required=True)
    parser.add_argument("--to", dest="to_date", required=True)
    parser.add_argument("--timezone", default="Asia/Kolkata")
    parser.add_argument("--output", required=True)
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO)
    result = run_report(
        from_date=date.fromisoformat(args.from_date),
        to_date=date.fromisoformat(args.to_date),
        tz_name=args.timezone,
        output_path=Path(args.output),
    )
    print(f"Export complete: {result['output']}")
    print(
        f"Cases: {result['cases']} | case-therapist pairs: {result['case_therapist_rows']} | "
        f"attended sessions: {result['attended_sessions']}"
    )


if __name__ == "__main__":
    main()
