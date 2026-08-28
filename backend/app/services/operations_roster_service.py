"""Operations roster exports — therapist and case Excel snapshots for admin Reports hub."""
from __future__ import annotations

import io
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime
from typing import Any, Optional
from zoneinfo import ZoneInfo

from sqlalchemy import extract, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.models.assignment import CaseAssignment, CaseAssignmentStatus
from app.models.attachment import Attachment
from app.models.case import Case, CaseStatus
from app.models.case_manager_meeting import CaseManagerMeeting, MeetingStatus
from app.models.child import Child
from app.models.client_billing import ClientInvoice
from app.models.daily_log import DailyLog
from app.models.incident import OPEN_INCIDENT_STATUSES, Incident
from app.models.ledger_billing import BillableStatus, BillingLedger
from app.models.parent import ParentGuardian, parent_child_link
from app.models.report import MonthlyReport, ObservationReport, ReportStatus
from app.models.session import Session as TherapySession
from app.models.session import SessionStatus
from app.models.support_ticket import SupportTicket, TicketStatus
from app.models.therapist_profile import TherapistProfile, TherapistProfileStatus
from app.models.user import User
from app.models.visibility import VisibilityStatus
from app.services.admin_iep_service import _derive_iep_status
from app.services.admin_scope_service import apply_case_scope, user_sees_global_cases
from app.services.billing_composer_service import normalize_billing_month
from app.services.user_provision_service import invite_status_for_email, login_ready

MAX_EXPORT_ROWS = 5000
IST = ZoneInfo("Asia/Kolkata")

CASE_COLUMNS: list[str] = [
    "case_id",
    "case_code",
    "child_id",
    "external_client_id",
    "client_name",
    "date_of_birth",
    "parent_name",
    "parent_email",
    "parent_phone",
    "parent_login_ready",
    "product_module",
    "service_type",
    "case_status",
    "operational_stage",
    "status_effective_date",
    "status_reason",
    "case_manager_user_id",
    "case_manager_name",
    "therapist_user_id",
    "therapist_name",
    "therapist_staff_id",
    "assignment_status",
    "assignment_start_date",
    "billing_type",
    "client_billing_mode",
    "client_rate_per_session_inr",
    "package_session_count",
    "package_amount_inr",
    "compensation_mode",
    "therapist_pay_inr",
    "sessions_scheduled",
    "sessions_completed",
    "sessions_cancelled",
    "sessions_rescheduled",
    "sessions_no_show",
    "sessions_client_absent",
    "billable_sessions",
    "missing_session_logs",
    "monthly_report_status",
    "monthly_report_parent_review",
    "monthly_report_missing",
    "observation_reports_count",
    "observation_latest_status",
    "iep_status",
    "meetings_scheduled",
    "meetings_completed",
    "last_meeting_type",
    "last_meeting_date",
    "open_tickets",
    "open_incidents",
    "invoice_generated",
    "invoice_amount_inr",
    "invoice_status",
]

THERAPIST_COLUMNS: list[str] = [
    "therapist_user_id",
    "external_employee_id",
    "therapist_profile_id",
    "full_name",
    "display_name",
    "email",
    "phone",
    "region",
    "employment_start_date",
    "employment_status",
    "account_status",
    "profile_status",
    "services_offered",
    "primary_cm_user_id",
    "primary_cm_name",
    "mentor_user_id",
    "mentor_name",
    "active_assignments_count",
    "active_case_ids",
    "active_case_codes",
    "case_managers_on_cases",
    "monthly_submitted",
    "monthly_approved",
    "monthly_under_review",
    "monthly_missing_cases",
    "sessions_completed",
    "sessions_cancelled",
    "sessions_rescheduled",
    "sessions_client_absent",
    "sessions_therapist_leave",
    "open_tickets_on_cases",
    "open_incidents_on_cases",
    "missing_session_logs_all_time",
    "module_assignments",
    "login_ready",
    "invite_status",
]

CASE_COLUMN_GUIDE: list[tuple[str, str, str]] = [
    ("case_id", "Internal case primary key", "Not editable — reference only"),
    ("case_code", "Human-readable case code", "Admin case detail"),
    ("child_id", "Client child record ID", "Admin client profiles"),
    ("external_client_id", "External client reference ID", "Admin client profiles / bulk import"),
    ("therapist_staff_id", "Therapist employee ID on active assignment", "Admin People → therapist ID"),
    ("case_manager_user_id", "Primary case manager user ID", "Admin case detail / allotment"),
    ("billing_type", "PER_SESSION or PACKAGE", "Admin case billing form"),
    ("client_billing_mode", "PREPAID or POSTPAID", "Admin case billing form"),
    ("monthly_report_missing", "Y if no monthly report for export month", "Therapist monthly reports"),
]

THERAPIST_COLUMN_GUIDE: list[tuple[str, str, str]] = [
    ("therapist_user_id", "Internal therapist user ID", "Not editable — reference only"),
    ("external_employee_id", "Staff / therapist ID", "Admin People → therapist ID"),
    ("primary_cm_user_id", "Primary case manager user ID", "Admin therapist profiles"),
    ("employment_start_date", "Employment start date", "Admin therapist profiles"),
    ("active_case_codes", "Semicolon-separated active case codes", "Case allotment"),
    ("monthly_missing_cases", "Active cases missing monthly report this month", "Reports hub"),
]


def default_export_month() -> str:
    return datetime.now(IST).strftime("%Y-%m")


def _month_long_label(ym: str) -> str:
    year_s, month_s = ym.split("-")[:2]
    return datetime(int(year_s), int(month_s), 1).strftime("%B %Y")


def _month_bounds(ym: str) -> tuple[date, date]:
    year_s, month_s = ym.split("-")[:2]
    y, m = int(year_s), int(month_s)
    return date(y, m, 1), date(y, m, monthrange(y, m)[1])


def _enum_val(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "value"):
        return str(value.value)
    return str(value)


def _account_status_label(user: User, db: Session) -> str:
    if not user.is_active:
        return "Deactivated"
    if login_ready(user, db):
        return "Active"
    if invite_status_for_email(db, user.email) == "pending":
        return "Invited"
    return "Inactive"


def _monthly_report_month_clause(ym: str):
    long_label = _month_long_label(ym)
    return or_(MonthlyReport.month.ilike(f"%{ym}%"), MonthlyReport.month.ilike(f"%{long_label}%"))


def _scoped_cases(
    db: Session,
    user: User,
    *,
    product_module: str | None = None,
) -> list[Case]:
    stmt = select(Case).options(selectinload(Case.child)).order_by(Case.case_code)
    stmt = apply_case_scope(stmt, user)
    if product_module:
        stmt = stmt.where(Case.product_module == product_module)
    return list(db.scalars(stmt).all())


def _parent_info_by_child(db: Session, child_ids: set[int]) -> dict[int, dict[str, Any]]:
    if not child_ids:
        return {}
    rows = db.execute(
        select(
            parent_child_link.c.child_id,
            User.full_name,
            User.email,
            User.phone,
            User.id,
            User.is_active,
            User.password_hash,
        )
        .join(ParentGuardian, ParentGuardian.id == parent_child_link.c.parent_guardian_id)
        .join(User, User.id == ParentGuardian.user_id)
        .where(parent_child_link.c.child_id.in_(child_ids))
        .order_by(parent_child_link.c.child_id, ParentGuardian.id)
    ).all()
    out: dict[int, dict[str, Any]] = {}
    for child_id, full_name, email, phone, user_id, is_active, password_hash in rows:
        if child_id in out:
            continue
        parent_user = User(id=user_id, email=email or "", is_active=is_active, password_hash=password_hash)
        out[child_id] = {
            "parent_name": full_name or "",
            "parent_email": email or "",
            "parent_phone": phone or "",
            "parent_login_ready": login_ready(parent_user, db),
        }
    return out


def _session_counts_by_case(db: Session, case_ids: list[int], ym: str) -> dict[int, dict[str, int]]:
    if not case_ids:
        return {}
    year_s, month_s = ym.split("-")[:2]
    y, m = int(year_s), int(month_s)
    counts: dict[int, dict[str, int]] = {cid: defaultdict(int) for cid in case_ids}

    rows = db.execute(
        select(TherapySession.case_id, TherapySession.status, func.count())
        .where(
            TherapySession.case_id.in_(case_ids),
            extract("year", TherapySession.scheduled_date) == y,
            extract("month", TherapySession.scheduled_date) == m,
        )
        .group_by(TherapySession.case_id, TherapySession.status)
    ).all()
    for case_id, status, cnt in rows:
        counts[case_id]["sessions_scheduled"] += int(cnt)
        key = {
            SessionStatus.COMPLETED: "sessions_completed",
            SessionStatus.CANCELLED: "sessions_cancelled",
            SessionStatus.RESCHEDULED: "sessions_rescheduled",
            SessionStatus.NO_SHOW: "sessions_no_show",
            SessionStatus.CLIENT_ABSENT: "sessions_client_absent",
        }.get(status)
        if key:
            counts[case_id][key] = int(cnt)

    billable_rows = db.execute(
        select(BillingLedger.case_id, func.count())
        .where(
            BillingLedger.case_id.in_(case_ids),
            BillingLedger.ledger_month == ym,
            BillingLedger.billable_status.in_([BillableStatus.BILLABLE, BillableStatus.INVOICED]),
        )
        .group_by(BillingLedger.case_id)
    ).all()
    for case_id, cnt in billable_rows:
        counts[case_id]["billable_sessions"] = int(cnt)

    return {k: dict(v) for k, v in counts.items()}


def _missing_logs_by_case(db: Session, case_ids: list[int]) -> dict[int, int]:
    if not case_ids:
        return {}
    return dict(
        db.execute(
            select(TherapySession.case_id, func.count())
            .outerjoin(DailyLog, DailyLog.session_id == TherapySession.id)
            .where(
                TherapySession.case_id.in_(case_ids),
                TherapySession.status == SessionStatus.COMPLETED,
                DailyLog.id.is_(None),
            )
            .group_by(TherapySession.case_id)
        ).all()
    )


def build_case_rows(
    db: Session,
    user: User,
    *,
    month: str | None = None,
    product_module: str | None = None,
) -> list[dict[str, Any]]:
    ym = normalize_billing_month(month or default_export_month())
    cases = _scoped_cases(db, user, product_module=product_module)
    if not cases:
        return []

    case_ids = [c.id for c in cases]
    child_ids = {c.child_id for c in cases}

    active_assignments = db.execute(
        select(
            CaseAssignment.case_id,
            CaseAssignment.therapist_user_id,
            CaseAssignment.status,
            CaseAssignment.start_date,
            User.full_name,
            User.external_employee_id,
        )
        .join(User, User.id == CaseAssignment.therapist_user_id)
        .where(
            CaseAssignment.case_id.in_(case_ids),
            CaseAssignment.status == CaseAssignmentStatus.ACTIVE,
        )
    ).all()
    assign_by_case: dict[int, tuple] = {}
    for row in active_assignments:
        assign_by_case[row.case_id] = row

    user_ids = {c.case_manager_user_id for c in cases if c.case_manager_user_id}
    user_ids.update(row.therapist_user_id for row in active_assignments)
    user_names: dict[int, str] = {}
    if user_ids:
        user_names = dict(db.execute(select(User.id, User.full_name).where(User.id.in_(user_ids))).all())

    parent_by_child = _parent_info_by_child(db, child_ids)
    session_counts = _session_counts_by_case(db, case_ids, ym)
    missing_logs = _missing_logs_by_case(db, case_ids)

    monthly_by_case: dict[int, MonthlyReport] = {}
    for report in db.scalars(
        select(MonthlyReport).where(
            MonthlyReport.case_id.in_(case_ids),
            _monthly_report_month_clause(ym),
            or_(MonthlyReport.category == "CLIENT_MONTHLY", MonthlyReport.category.is_(None)),
        )
    ).all():
        prev = monthly_by_case.get(report.case_id)
        if not prev or (report.updated_at or report.created_at) > (prev.updated_at or prev.created_at):
            monthly_by_case[report.case_id] = report

    obs_count: dict[int, int] = dict(
        db.execute(
            select(ObservationReport.case_id, func.count())
            .where(
                ObservationReport.case_id.in_(case_ids),
                extract("year", ObservationReport.created_at) == int(ym.split("-")[0]),
                extract("month", ObservationReport.created_at) == int(ym.split("-")[1]),
            )
            .group_by(ObservationReport.case_id)
        ).all()
    )
    obs_latest: dict[int, str] = {}
    for report in db.scalars(
        select(ObservationReport)
        .where(ObservationReport.case_id.in_(case_ids))
        .order_by(ObservationReport.case_id, ObservationReport.created_at.desc(), ObservationReport.id.desc())
    ).all():
        if report.case_id not in obs_latest:
            obs_latest[report.case_id] = _enum_val(report.status)

    iep_atts = db.scalars(
        select(Attachment).where(Attachment.case_id.in_(case_ids), Attachment.entity_type == "iep")
    ).all()
    iep_by_case: dict[int, Attachment] = {}
    for att in iep_atts:
        prev = iep_by_case.get(att.case_id)
        if not prev or att.created_at > prev.created_at:
            iep_by_case[att.case_id] = att

    start_d, end_d = _month_bounds(ym)
    meetings_scheduled = dict(
        db.execute(
            select(CaseManagerMeeting.case_id, func.count())
            .where(
                CaseManagerMeeting.case_id.in_(case_ids),
                CaseManagerMeeting.scheduled_date >= start_d,
                CaseManagerMeeting.scheduled_date <= end_d,
            )
            .group_by(CaseManagerMeeting.case_id)
        ).all()
    )
    meetings_completed = dict(
        db.execute(
            select(CaseManagerMeeting.case_id, func.count())
            .where(
                CaseManagerMeeting.case_id.in_(case_ids),
                CaseManagerMeeting.scheduled_date >= start_d,
                CaseManagerMeeting.scheduled_date <= end_d,
                CaseManagerMeeting.status == MeetingStatus.COMPLETED,
            )
            .group_by(CaseManagerMeeting.case_id)
        ).all()
    )
    last_meeting_by_case: dict[int, CaseManagerMeeting] = {}
    for meeting in db.scalars(
        select(CaseManagerMeeting)
        .where(CaseManagerMeeting.case_id.in_(case_ids))
        .order_by(CaseManagerMeeting.case_id, CaseManagerMeeting.scheduled_date.desc())
    ).all():
        if meeting.case_id and meeting.case_id not in last_meeting_by_case:
            last_meeting_by_case[meeting.case_id] = meeting

    ticket_counts = dict(
        db.execute(
            select(SupportTicket.case_id, func.count())
            .where(
                SupportTicket.case_id.in_(case_ids),
                SupportTicket.status.in_([TicketStatus.OPEN, TicketStatus.IN_PROGRESS]),
            )
            .group_by(SupportTicket.case_id)
        ).all()
    )
    incident_counts = dict(
        db.execute(
            select(Incident.case_id, func.count())
            .where(
                Incident.case_id.in_(case_ids),
                Incident.status.in_(list(OPEN_INCIDENT_STATUSES)),
            )
            .group_by(Incident.case_id)
        ).all()
    )

    invoices_by_case: dict[int, ClientInvoice] = {}
    for inv in db.scalars(
        select(ClientInvoice).where(
            ClientInvoice.case_id.in_(case_ids),
            ClientInvoice.billing_month == ym,
        )
    ).all():
        invoices_by_case[inv.case_id] = inv

    submitted_monthly_case_ids = set(monthly_by_case.keys())
    rows: list[dict[str, Any]] = []
    for case in cases:
        child = case.child
        parent = parent_by_child.get(case.child_id, {})
        assign = assign_by_case.get(case.id)
        therapist_user_id = assign.therapist_user_id if assign else None
        therapist_name = assign.full_name if assign else ""
        therapist_staff_id = assign.external_employee_id if assign else ""
        assignment_status = _enum_val(assign.status) if assign else ""
        assignment_start = assign.start_date.isoformat() if assign and assign.start_date else ""

        cm_id = case.case_manager_user_id
        sess = session_counts.get(case.id, {})
        monthly = monthly_by_case.get(case.id)
        missing_monthly = (
            case.status == CaseStatus.ACTIVE and case.id not in submitted_monthly_case_ids
        )
        inv = invoices_by_case.get(case.id)
        pay_inr = ""
        if case.compensation_mode and case.compensation_mode.value == "FIXED_LUMP":
            pay_inr = case.therapist_fixed_pay_inr
        elif case.pay_share_amount_inr is not None:
            pay_inr = case.pay_share_amount_inr

        last_meeting = last_meeting_by_case.get(case.id)
        rows.append(
            {
                "case_id": case.id,
                "case_code": case.case_code,
                "child_id": case.child_id,
                "external_client_id": child.external_client_id if child else "",
                "client_name": child.full_name if child else "",
                "date_of_birth": child.date_of_birth.isoformat() if child and child.date_of_birth else "",
                "parent_name": parent.get("parent_name", ""),
                "parent_email": parent.get("parent_email", ""),
                "parent_phone": parent.get("parent_phone", ""),
                "parent_login_ready": parent.get("parent_login_ready", False),
                "product_module": case.product_module,
                "service_type": case.service_type,
                "case_status": _enum_val(case.status),
                "operational_stage": case.operational_stage or "",
                "status_effective_date": case.status_effective_date.isoformat() if case.status_effective_date else "",
                "status_reason": (case.status_reason or "")[:500],
                "case_manager_user_id": cm_id or "",
                "case_manager_name": user_names.get(cm_id, "") if cm_id else "",
                "therapist_user_id": therapist_user_id or "",
                "therapist_name": therapist_name or "",
                "therapist_staff_id": therapist_staff_id or "",
                "assignment_status": assignment_status,
                "assignment_start_date": assignment_start,
                "billing_type": _enum_val(case.billing_type),
                "client_billing_mode": _enum_val(case.client_billing_mode),
                "client_rate_per_session_inr": case.client_rate_per_session_inr or "",
                "package_session_count": case.package_session_count or "",
                "package_amount_inr": case.package_amount_inr or "",
                "compensation_mode": _enum_val(case.compensation_mode),
                "therapist_pay_inr": pay_inr,
                "sessions_scheduled": sess.get("sessions_scheduled", 0),
                "sessions_completed": sess.get("sessions_completed", 0),
                "sessions_cancelled": sess.get("sessions_cancelled", 0),
                "sessions_rescheduled": sess.get("sessions_rescheduled", 0),
                "sessions_no_show": sess.get("sessions_no_show", 0),
                "sessions_client_absent": sess.get("sessions_client_absent", 0),
                "billable_sessions": sess.get("billable_sessions", 0),
                "missing_session_logs": int(missing_logs.get(case.id, 0)),
                "monthly_report_status": _enum_val(monthly.status) if monthly else "",
                "monthly_report_parent_review": monthly.parent_review_status or "" if monthly else "",
                "monthly_report_missing": "Y" if missing_monthly else "N",
                "observation_reports_count": int(obs_count.get(case.id, 0)),
                "observation_latest_status": obs_latest.get(case.id, ""),
                "iep_status": _derive_iep_status(iep_by_case.get(case.id)),
                "meetings_scheduled": int(meetings_scheduled.get(case.id, 0)),
                "meetings_completed": int(meetings_completed.get(case.id, 0)),
                "last_meeting_type": _enum_val(last_meeting.meeting_type) if last_meeting else "",
                "last_meeting_date": last_meeting.scheduled_date.isoformat() if last_meeting else "",
                "open_tickets": int(ticket_counts.get(case.id, 0)),
                "open_incidents": int(incident_counts.get(case.id, 0)),
                "invoice_generated": "Y" if inv else "N",
                "invoice_amount_inr": float(inv.total_inr) if inv else "",
                "invoice_status": _enum_val(inv.status) if inv else "",
            }
        )
    return rows


def build_therapist_rows(
    db: Session,
    user: User,
    *,
    month: str | None = None,
    product_module: str | None = None,
) -> list[dict[str, Any]]:
    ym = normalize_billing_month(month or default_export_month())
    cases = _scoped_cases(db, user, product_module=product_module)
    scoped_case_ids = {c.id for c in cases}

    if user_sees_global_cases(user):
        profiles = list(
            db.scalars(
                select(TherapistProfile)
                .options(selectinload(TherapistProfile.user))
                .where(TherapistProfile.status != TherapistProfileStatus.DELETED)
                .order_by(TherapistProfile.id)
            ).all()
        )
    else:
        therapist_ids = {
            row.therapist_user_id
            for row in db.execute(
                select(CaseAssignment.therapist_user_id).where(
                    CaseAssignment.case_id.in_(scoped_case_ids or [-1]),
                    CaseAssignment.status == CaseAssignmentStatus.ACTIVE,
                )
            ).all()
        }
        profiles = list(
            db.scalars(
                select(TherapistProfile)
                .options(selectinload(TherapistProfile.user))
                .where(TherapistProfile.user_id.in_(therapist_ids or [-1]))
            ).all()
        )

    if not profiles:
        return []

    therapist_ids = [p.user_id for p in profiles]
    supervisor_ids = {p.supervisor_user_id for p in profiles if p.supervisor_user_id}
    mentor_ids = {p.mentor_user_id for p in profiles if p.mentor_user_id}
    name_ids = supervisor_ids | mentor_ids
    staff_names: dict[int, str] = {}
    if name_ids:
        staff_names = dict(db.execute(select(User.id, User.full_name).where(User.id.in_(name_ids))).all())

    assignments = db.execute(
        select(
            CaseAssignment.therapist_user_id,
            CaseAssignment.case_id,
            Case.case_code,
            Case.case_manager_user_id,
            Case.status,
        )
        .join(Case, Case.id == CaseAssignment.case_id)
        .where(
            CaseAssignment.therapist_user_id.in_(therapist_ids),
            CaseAssignment.status == CaseAssignmentStatus.ACTIVE,
            CaseAssignment.case_id.in_(scoped_case_ids or [-1]),
        )
    ).all()
    cases_by_therapist: dict[int, list] = defaultdict(list)
    cm_names_by_therapist: dict[int, set[str]] = defaultdict(set)
    cm_ids_for_names: set[int] = set()
    for t_id, case_id, case_code, cm_id, case_status in assignments:
        cases_by_therapist[t_id].append((case_id, case_code, case_status))
        if cm_id:
            cm_ids_for_names.add(cm_id)
    cm_name_map = dict(
        db.execute(select(User.id, User.full_name).where(User.id.in_(cm_ids_for_names))).all()
    ) if cm_ids_for_names else {}
    for t_id, case_id, case_code, cm_id, _ in assignments:
        if cm_id and cm_id in cm_name_map:
            cm_names_by_therapist[t_id].add(cm_name_map[cm_id])

    year_s, month_s = ym.split("-")[:2]
    y, m = int(year_s), int(month_s)

    def _session_agg(status: SessionStatus) -> dict[int, int]:
        return dict(
            db.execute(
                select(TherapySession.therapist_user_id, func.count())
                .where(
                    TherapySession.therapist_user_id.in_(therapist_ids),
                    TherapySession.case_id.in_(scoped_case_ids or [-1]),
                    extract("year", TherapySession.scheduled_date) == y,
                    extract("month", TherapySession.scheduled_date) == m,
                    TherapySession.status == status,
                )
                .group_by(TherapySession.therapist_user_id)
            ).all()
        )

    completed_sess = _session_agg(SessionStatus.COMPLETED)
    cancelled_sess = _session_agg(SessionStatus.CANCELLED)
    rescheduled_sess = _session_agg(SessionStatus.RESCHEDULED)
    client_absent_sess = _session_agg(SessionStatus.CLIENT_ABSENT)
    therapist_leave_sess = _session_agg(SessionStatus.THERAPIST_LEAVE)

    monthly_reports = db.scalars(
        select(MonthlyReport).where(
            MonthlyReport.therapist_user_id.in_(therapist_ids),
            MonthlyReport.case_id.in_(scoped_case_ids or [-1]),
            _monthly_report_month_clause(ym),
        )
    ).all()
    monthly_stats: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    monthly_cases_by_therapist: dict[int, set[int]] = defaultdict(set)
    for report in monthly_reports:
        monthly_stats[report.therapist_user_id]["monthly_submitted"] += 1
        monthly_cases_by_therapist[report.therapist_user_id].add(report.case_id)
        if report.status == ReportStatus.APPROVED or report.status == ReportStatus.PUBLISHED:
            monthly_stats[report.therapist_user_id]["monthly_approved"] += 1
        if report.status == ReportStatus.UNDER_REVIEW:
            monthly_stats[report.therapist_user_id]["monthly_under_review"] += 1

    missing_monthly_by_therapist: dict[int, int] = defaultdict(int)
    for t_id, case_rows in cases_by_therapist.items():
        submitted_cases = monthly_cases_by_therapist.get(t_id, set())
        for case_id, _, case_status in case_rows:
            if case_status == CaseStatus.ACTIVE and case_id not in submitted_cases:
                missing_monthly_by_therapist[t_id] += 1

    therapist_case_ids = list(scoped_case_ids) if scoped_case_ids else [-1]
    missing_logs_by_case = _missing_logs_by_case(db, therapist_case_ids)
    case_to_therapist: dict[int, int] = {}
    for t_id, case_rows in cases_by_therapist.items():
        for case_id, _, _ in case_rows:
            case_to_therapist[case_id] = t_id
    missing_logs_by_therapist: dict[int, int] = defaultdict(int)
    for case_id, cnt in missing_logs_by_case.items():
        t_id = case_to_therapist.get(case_id)
        if t_id:
            missing_logs_by_therapist[t_id] += int(cnt)

    ticket_by_case = dict(
        db.execute(
            select(SupportTicket.case_id, func.count())
            .where(
                SupportTicket.case_id.in_(therapist_case_ids),
                SupportTicket.status.in_([TicketStatus.OPEN, TicketStatus.IN_PROGRESS]),
            )
            .group_by(SupportTicket.case_id)
        ).all()
    )
    incident_by_case = dict(
        db.execute(
            select(Incident.case_id, func.count())
            .where(
                Incident.case_id.in_(therapist_case_ids),
                Incident.status.in_(list(OPEN_INCIDENT_STATUSES)),
            )
            .group_by(Incident.case_id)
        ).all()
    )
    tickets_by_therapist: dict[int, int] = defaultdict(int)
    incidents_by_therapist: dict[int, int] = defaultdict(int)
    for case_id, cnt in ticket_by_case.items():
        t_id = case_to_therapist.get(case_id)
        if t_id:
            tickets_by_therapist[t_id] += int(cnt)
    for case_id, cnt in incident_by_case.items():
        t_id = case_to_therapist.get(case_id)
        if t_id:
            incidents_by_therapist[t_id] += int(cnt)

    rows: list[dict[str, Any]] = []
    for profile in profiles:
        u = profile.user
        if not u:
            continue
        t_id = profile.user_id
        case_rows = cases_by_therapist.get(t_id, [])
        case_ids = [str(c[0]) for c in case_rows]
        case_codes = [c[1] for c in case_rows]
        stats = monthly_stats.get(t_id, {})
        services = profile.services_offered or []
        if isinstance(services, list):
            services_str = ", ".join(str(s) for s in services)
        else:
            services_str = str(services)
        rows.append(
            {
                "therapist_user_id": t_id,
                "external_employee_id": u.external_employee_id or "",
                "therapist_profile_id": profile.id,
                "full_name": u.full_name or "",
                "display_name": profile.display_name or "",
                "email": u.email or "",
                "phone": u.phone or "",
                "region": u.region or "",
                "employment_start_date": profile.employment_start_date.isoformat()
                if profile.employment_start_date
                else "",
                "employment_status": _enum_val(u.employment_status),
                "account_status": _account_status_label(u, db),
                "profile_status": _enum_val(profile.status),
                "services_offered": services_str,
                "primary_cm_user_id": profile.supervisor_user_id or "",
                "primary_cm_name": staff_names.get(profile.supervisor_user_id, "")
                if profile.supervisor_user_id
                else "",
                "mentor_user_id": profile.mentor_user_id or "",
                "mentor_name": staff_names.get(profile.mentor_user_id, "") if profile.mentor_user_id else "",
                "active_assignments_count": len(case_rows),
                "active_case_ids": ";".join(case_ids),
                "active_case_codes": ";".join(case_codes),
                "case_managers_on_cases": ";".join(sorted(cm_names_by_therapist.get(t_id, set()))),
                "monthly_submitted": int(stats.get("monthly_submitted", 0)),
                "monthly_approved": int(stats.get("monthly_approved", 0)),
                "monthly_under_review": int(stats.get("monthly_under_review", 0)),
                "monthly_missing_cases": int(missing_monthly_by_therapist.get(t_id, 0)),
                "sessions_completed": int(completed_sess.get(t_id, 0)),
                "sessions_cancelled": int(cancelled_sess.get(t_id, 0)),
                "sessions_rescheduled": int(rescheduled_sess.get(t_id, 0)),
                "sessions_client_absent": int(client_absent_sess.get(t_id, 0)),
                "sessions_therapist_leave": int(therapist_leave_sess.get(t_id, 0)),
                "open_tickets_on_cases": int(tickets_by_therapist.get(t_id, 0)),
                "open_incidents_on_cases": int(incidents_by_therapist.get(t_id, 0)),
                "missing_session_logs_all_time": int(missing_logs_by_therapist.get(t_id, 0)),
                "module_assignments": ", ".join(u.module_assignments or []),
                "login_ready": login_ready(u, db),
                "invite_status": invite_status_for_email(db, u.email),
            }
        )
    return rows


def _workbook_bytes(
    *,
    title: str,
    subtitle: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    column_guide: list[tuple[str, str, str]],
    user: User,
    filters: dict[str, str],
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    from app.services.export_document_service import export_meta, xlsx_footer_rows, xlsx_preamble_rows

    meta = export_meta(user)
    filter_line = " · ".join(f"{k}={v}" for k, v in filters.items() if v)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    preamble = xlsx_preamble_rows(title, subtitle, meta)
    if filter_line:
        preamble.insert(3, [f"Filters: {filter_line}"])
    for row in preamble:
        ws.append(row)
    header_row_idx = len(preamble) + 1
    ws.append(columns)
    for item in rows:
        ws.append([item.get(col, "") for col in columns])
    for row in xlsx_footer_rows(meta):
        ws.append(row)
    if rows:
        ws.auto_filter.ref = f"A{header_row_idx}:{get_column_letter(len(columns))}{header_row_idx + len(rows)}"
        ws.freeze_panes = f"A{header_row_idx + 1}"

    guide = wb.create_sheet("Column guide")
    guide.append(["Field", "Description", "Editable via"])
    for field, desc, editable in column_guide:
        guide.append([field, desc, editable])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_case_roster_xlsx(
    db: Session,
    user: User,
    *,
    month: str | None = None,
    product_module: str | None = None,
) -> bytes:
    ym = normalize_billing_month(month or default_export_month())
    rows = build_case_rows(db, user, month=ym, product_module=product_module)
    if len(rows) > MAX_EXPORT_ROWS:
        raise ValueError(
            f"Export has {len(rows)} rows (max {MAX_EXPORT_ROWS}). Narrow product_module filter."
        )
    return _workbook_bytes(
        title="Case operations roster",
        subtitle="One row per case — client, billing, sessions, reports, IEP, meetings",
        columns=CASE_COLUMNS,
        rows=rows,
        column_guide=CASE_COLUMN_GUIDE,
        user=user,
        filters={"month": ym, "product_module": product_module or "all"},
    )


def export_therapist_roster_xlsx(
    db: Session,
    user: User,
    *,
    month: str | None = None,
    product_module: str | None = None,
) -> bytes:
    ym = normalize_billing_month(month or default_export_month())
    rows = build_therapist_rows(db, user, month=ym, product_module=product_module)
    if len(rows) > MAX_EXPORT_ROWS:
        raise ValueError(
            f"Export has {len(rows)} rows (max {MAX_EXPORT_ROWS}). Narrow product_module filter."
        )
    return _workbook_bytes(
        title="Therapist operations roster",
        subtitle="One row per therapist — People module, caseload, reports, sessions",
        columns=THERAPIST_COLUMNS,
        rows=rows,
        column_guide=THERAPIST_COLUMN_GUIDE,
        user=user,
        filters={"month": ym, "product_module": product_module or "all"},
    )
