from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.core.database import SessionLocal
from app.main import app
from app.models.support_ticket import SupportTicket, TicketCategory, TicketMessage, TicketStatus
from app.models.user import User
from app.seed.demo_seed import run as seed_run

client = TestClient(app)

SECRET = "ZXCHILDNAME9k2"
BANK = "ZXBANK998877"
CLINICAL = "ZXCLINICALNOTE9k2"
RANGE = {"date_from": "2020-01-01", "date_to": "2020-01-31"}


@pytest.fixture(scope="module", autouse=True)
def setup_db():
    seed_run()


def _auth_headers(email: str = "superadmin@demo.com"):
    response = client.post("/api/v1/auth/login", json={"email": email, "password": "demo123"})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _create_ticket(headers, *, subject, body, category):
    response = client.post(
        "/api/v1/tickets",
        headers=headers,
        json={"subject": subject, "body": body, "category": category},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _stamp(ticket_id: int, opened: datetime, *, status: TicketStatus, module: str | None, assignee_email: str | None = None):
    with SessionLocal() as db:
        ticket = db.get(SupportTicket, ticket_id)
        assert ticket is not None
        ticket.created_at = opened
        ticket.status = status
        ticket.product_module = module
        if assignee_email:
            assignee = db.query(User).filter(User.email == assignee_email).one()
            ticket.assigned_to_user_id = assignee.id
        else:
            ticket.assigned_to_user_id = None
        for message in db.query(TicketMessage).filter(TicketMessage.ticket_id == ticket_id).all():
            message.created_at = opened
        db.commit()


def _reply(ticket_id: int, author_email: str, opened: datetime, hours_later: float, body: str):
    with SessionLocal() as db:
        author = db.query(User).filter(User.email == author_email).one()
        db.add(
            TicketMessage(
                ticket_id=ticket_id,
                author_user_id=author.id,
                body=body,
                created_at=opened + timedelta(hours=hours_later),
            )
        )
        db.commit()
        return author.full_name


@pytest.fixture(scope="module")
def january_tickets():
    therapist = _auth_headers("therapist@demo.com")
    opened = datetime(2020, 1, 15, 4, 0, tzinfo=timezone.utc)
    pay_id = _create_ticket(
        therapist,
        subject="Salary discrepancy",
        body=f"The payout does not match. {SECRET} {BANK}",
        category="FINANCE",
    )
    _stamp(pay_id, opened, status=TicketStatus.OPEN, module="homecare", assignee_email="finance@demo.com")
    finance_name = _reply(pay_id, "finance@demo.com", opened, 10, "Payout status is paid.")

    leave_opened = datetime(2020, 1, 2, 4, 0, tzinfo=timezone.utc)
    leave_id = _create_ticket(
        therapist,
        subject="Query regarding approved leave",
        body="Leave balance is not showing",
        category="HR",
    )
    _stamp(leave_id, leave_opened, status=TicketStatus.IN_PROGRESS, module="shadow_support")

    other_id = _create_ticket(
        therapist,
        subject="Hello there",
        body="A general question",
        category="OTHER",
    )
    _stamp(other_id, datetime(2020, 1, 20, 4, 0, tzinfo=timezone.utc), status=TicketStatus.OPEN, module=None)

    posh_id = _create_ticket(
        therapist,
        subject="Incident at the session",
        body=f"Please review {CLINICAL}",
        category="POSH",
    )
    _stamp(posh_id, datetime(2020, 1, 18, 4, 0, tzinfo=timezone.utc), status=TicketStatus.OPEN, module="b2b")

    return {
        "pay_id": pay_id,
        "leave_id": leave_id,
        "other_id": other_id,
        "posh_id": posh_id,
        "finance_name": finance_name,
    }


def _report(headers, **params):
    response = client.get("/api/v1/admin/support/ticket-report", headers=headers, params={**RANGE, **params})
    assert response.status_code == 200, response.text
    return response.json()


def test_therapist_cannot_open_the_report():
    response = client.get("/api/v1/admin/support/ticket-report", headers=_auth_headers("therapist@demo.com"))
    assert response.status_code == 403


def test_report_does_not_write_tickets(january_tickets):
    headers = _auth_headers()
    with SessionLocal() as db:
        before_messages = db.query(TicketMessage).count()
        before_status = db.get(SupportTicket, january_tickets["pay_id"]).status
    response = client.post("/api/v1/admin/support/ticket-report", headers=headers, json={"body": "auto reply"})
    assert response.status_code == 405
    _report(headers)
    with SessionLocal() as db:
        assert db.query(TicketMessage).count() == before_messages
        assert db.get(SupportTicket, january_tickets["pay_id"]).status == before_status


def test_report_tables_filters_and_privacy(january_tickets):
    headers = _auth_headers()
    payload = _report(headers)
    assert payload["read_only"] is True
    assert payload["filters"]["date_from"] == "2020-01-01"
    assert payload["filters"]["timezone"] == "Asia/Kolkata"
    assert payload["total"] >= 4

    by_status = {row["status"]: row["tickets"] for row in payload["by_status"]}
    assert by_status["OPEN"] >= 3
    assert by_status["IN_PROGRESS"] >= 1
    assert sum(row["tickets"] for row in payload["by_status"]) == payload["total"]
    assert payload["in_flight"] == by_status["OPEN"] + by_status["IN_PROGRESS"]

    categories = {row["category"]: row for row in payload["by_category"]}
    assert categories["FINANCE"]["total"] >= 1
    assert categories["HR"]["in_progress"] >= 1
    assert categories["POSH"]["counts_only"] is True
    assert categories["CPP"]["counts_only"] is True

    modules = {row["module"]: row for row in payload["by_module"]}
    assert set(modules) >= {"shadow_support", "homecare", "b2b", "none"}
    assert modules["homecare"]["total"] >= 1
    assert modules["none"]["total"] >= 1
    assert modules["b2b"]["total"] >= 1

    roles = {row["role"]: row["tickets"] for row in payload["raised_by_role"]}
    assert roles.get("THERAPIST", 0) >= 4

    finance_replies = [row for row in payload["staff_replies"] if row["name"] == january_tickets["finance_name"]]
    assert finance_replies
    assert finance_replies[0]["replies"] >= 1
    assert finance_replies[0]["first_replies"] >= 1
    assert "Finance" in finance_replies[0]["roles"]

    hours = {row["category"]: row for row in payload["first_reply_hours"]}
    assert hours["FINANCE"]["median_hours"] == 10.0
    assert hours["FINANCE"]["tickets_with_reply"] >= 1
    assert "suggested_reply" not in hours["POSH"]
    assert "suggested_reply" not in hours["CPP"]

    groups = {row["key"]: row for row in payload["question_groups"]}
    assert groups["pay_mismatch"]["total"] >= 1
    assert groups["leave"]["still_open"] >= 1
    assert groups["clinical_or_incident"]["total"] >= 1
    assert groups["other"]["total"] >= 1
    assert sum(row["total"] for row in payload["question_groups"]) == payload["total"]
    assert all("suggested" not in row for row in payload["question_groups"])

    queue_ids = {row["id"] for row in payload["queue"]}
    assert january_tickets["pay_id"] in queue_ids
    assert january_tickets["leave_id"] in queue_ids
    assert january_tickets["posh_id"] not in queue_ids
    pay_row = next(row for row in payload["queue"] if row["id"] == january_tickets["pay_id"])
    assert pay_row["awaiting_first_reply"] is False
    assert pay_row["module"] == "homecare"
    assert pay_row["assignee"] == january_tickets["finance_name"]
    assert "subject" not in pay_row
    assert "body" not in pay_row
    other_row = next(row for row in payload["queue"] if row["id"] == january_tickets["other_id"])
    assert other_row["awaiting_first_reply"] is True
    assert other_row["case_code"] in (None, "")
    assert payload["aging"]["open_no_reply"] >= 1
    assert payload["aging"]["in_progress_older_than_7_days"] >= 1
    assert payload["aging"]["in_progress_older_than_30_days"] >= 1
    assert payload["restricted_omitted"] >= 1

    blob = response_text(payload)
    assert SECRET not in blob
    assert BANK not in blob
    assert CLINICAL not in blob

    homecare = _report(headers, product_module="homecare")
    assert homecare["total"] >= 1
    assert all(row["module"] == "homecare" for row in homecare["by_module"] if row["total"])
    assert january_tickets["pay_id"] in {row["id"] for row in homecare["queue"]}
    assert january_tickets["other_id"] not in {row["id"] for row in homecare["queue"]}

    finance_only = _report(headers, category="FINANCE", status="OPEN")
    assert finance_only["total"] >= 1
    assert all(row["category"] != "HR" or row["total"] == 0 for row in finance_only["by_category"])
    assert {row["id"] for row in finance_only["queue"]} >= {january_tickets["pay_id"]}
    assert january_tickets["leave_id"] not in {row["id"] for row in finance_only["queue"]}

    with SessionLocal() as db:
        finance = db.query(User).filter(User.email == "finance@demo.com").one()
        finance_id = finance.id
    assigned = _report(headers, assigned_to=str(finance_id))
    assert january_tickets["pay_id"] in {row["id"] for row in assigned["queue"]}
    assert january_tickets["other_id"] not in {row["id"] for row in assigned["queue"]}

    finance_headers = _auth_headers("finance@demo.com")
    finance_view = _report(finance_headers)
    finance_ids = {row["id"] for row in finance_view["queue"]}
    assert january_tickets["pay_id"] in finance_ids
    assert january_tickets["other_id"] not in finance_ids
    assert SECRET not in response_text(finance_view)


def test_excel_export_matches_the_tables(january_tickets):
    headers = _auth_headers()
    response = client.get("/api/v1/admin/support/ticket-report.xlsx", headers=headers, params=RANGE)
    assert response.status_code == 200, response.text
    assert "spreadsheet" in response.headers["content-type"]
    assert "support-tickets-report-2020-01-01-to-2020-01-31.xlsx" in response.headers["content-disposition"]
    text = response.content.decode("latin1")
    assert SECRET not in text
    assert BANK not in text
    assert CLINICAL not in text
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == [
        "Status",
        "Live queue",
        "By category",
        "By module",
        "Raised by role",
        "Staff replies",
        "First reply hours",
        "Aging",
        "Question groups",
    ]
    queue = workbook["Live queue"]
    header_index = None
    ticket_ids = []
    for index, row in enumerate(queue.iter_rows(values_only=True)):
        if row and row[0] == "Ticket":
            header_index = index
            assert "Case code" in row
            assert "Subject" not in row
            continue
        if header_index is None or not row or row[0] in (None, ""):
            continue
        ticket_ids.append(row[0])
    assert header_index is not None
    assert january_tickets["pay_id"] in ticket_ids
    assert january_tickets["posh_id"] not in ticket_ids
    assert "suggested" not in " ".join(workbook.sheetnames).lower()


def response_text(payload) -> str:
    return str(payload)
